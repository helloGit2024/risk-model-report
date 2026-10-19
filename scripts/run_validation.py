# -*- coding: utf-8 -*-
"""
端到端验证脚本
================================================================================
两部分验证，互补覆盖：

A. 真实数据回归（data/ 51.8万行）
   —— 核心是 KS/AUC 必须对齐实测黄金基线，确保重构未改变风控口径
B. Mock 异常分支（mock/ 5万行）
   —— 覆盖真实数据触达不到的边界场景，确保不崩溃

运行：python3 validate_report.py
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import tempfile
import time
import zipfile

import numpy as np
import pandas as pd

from riskpkg.metrics import RiskMetrics
from riskpkg.stability import StabilityMonitor

# ---------------------------------------------------------------------------
# 真实数据黄金基线（实测，作为重构回归的判定标准）
# ---------------------------------------------------------------------------
GOLDEN = {
    "train": {"ks": 0.2656, "auc": 0.6864},
    "test": {"ks": 0.2063, "auc": 0.6474},
    "oot": {"ks": 0.2191, "auc": 0.6510},
}
TOL = 1e-3

RESULTS = []


def check(name: str, ok: bool, detail: str = "") -> bool:
    status = "PASS" if ok else "FAIL"
    RESULTS.append((status, name, detail))
    print(f"  [{status}] {name}" + (f"  {detail}" if detail else ""))
    return ok


def sheet_layout(xlsx: str, sheet: str):
    """
    读取单 Sheet 版式：(空行列表, [(类型, 起行, 止行, 起列序号)], worksheet)。

    用于定版校验：分组空行是否存在、条件格式是否跨越空行。
    """
    import re
    import openpyxl
    from openpyxl.utils import column_index_from_string as col_idx

    wb = openpyxl.load_workbook(xlsx)
    ws = wb[sheet]
    empties = [i for i, row in enumerate(ws.iter_rows(values_only=True), 1)
               if all(v is None or v == "" for v in row)]
    rngs = []
    for entry in ws.conditional_formatting:
        t = entry.rules[0].type if entry.rules else ""
        if t not in ("dataBar", "colorScale"):
            continue
        for area in entry.sqref:
            m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", str(area))
            if m:
                rngs.append((t, int(m.group(2)), int(m.group(4)), col_idx(m.group(1))))
    return empties, rngs, ws


def count_conditional_formats(xlsx: str) -> dict:
    """统计 xlsx 内的条件格式数量（数据条/色阶/图标集），验证美化是否真正写入。"""
    out = {"dataBar": 0, "colorScale": 0, "iconSet": 0}
    if not os.path.exists(xlsx):
        return out
    with zipfile.ZipFile(xlsx) as z:
        for n in z.namelist():
            if n.startswith("xl/worksheets/sheet") and n.endswith(".xml"):
                x = z.read(n).decode("utf8")
                out["dataBar"] += x.count("dataBar")
                out["colorScale"] += x.count("colorScale")
                out["iconSet"] += x.count("iconSet")
    return out


# ===========================================================================
# A. 真实数据验证
# ===========================================================================
def validate_real(data_dir: str = None) -> None:
    print("\n" + "=" * 78)
    print("A. 真实数据验证（可选：--real-data <dir>；不提供则跳过）")
    print("=" * 78)

    # 真实数据涉敏，不随 skill 分发；有数据的人才需要跑这段回归
    data_dir = data_dir or os.environ.get("RISK_REAL_DATA") or "data"
    cfg_path = os.path.join(data_dir, "main_config.json")
    if not os.path.exists(cfg_path):
        print(f"  SKIP：未找到 {cfg_path}。真实数据不随 skill 分发（属正常）；"
              f"如需回归请 --real-data <dir>")
        return

    import generate_report

    cfg = json.load(open(cfg_path))
    report_path = os.path.join(tempfile.mkdtemp(prefix="riskreport_"), "report.xlsx")
    score_col = "D05_lgb_pkl"

    t0 = time.time()
    generate_report.get_report(cfg, score_col, os.path.join(data_dir, "D05_newdf.csv"),
                               os.path.join(data_dir, "lgbdoc.csv"),
                               report_path,
                               os.path.join(data_dir, "D05_lgb_importance.csv"), top=10)
    elapsed = time.time() - t0
    check("报告生成无异常", os.path.exists(report_path), f"耗时 {elapsed:.1f}s")

    # --- Sheet 完整性 ---
    sheets = pd.ExcelFile(report_path).sheet_names
    check("Sheet 数量 >= 15", len(sheets) >= 15, f"实际 {len(sheets)} 个")
    # 定版清单（逐月PSI / 数据集PSI 等已按期望版删除，不再作为必需项）
    required = ["建模效果统计", "模型分10箱", "PSI热力矩阵", "特征全景",
                "单变量分析", "单变量分箱按月分析", "校准度"]
    missing = [s for s in required if s not in sheets]
    check("关键 Sheet 齐全", not missing, f"缺失 {missing}" if missing else str(required))

    # --- 核心指标回归（对齐黄金基线）---
    mm = pd.read_excel(report_path, "建模效果统计", header=1)
    mm = mm.set_index(mm.columns[0])
    for flag, exp in GOLDEN.items():
        if flag not in mm.index:
            check(f"{flag} 指标存在", False, "缺失该数据集行")
            continue
        ks = float(mm.loc[flag, "ks"])
        auc = float(mm.loc[flag, "auc"])
        check(f"{flag} KS 对齐基线", abs(ks - exp["ks"]) < TOL,
              f"{ks:.4f} vs {exp['ks']}")
        check(f"{flag} AUC 对齐基线", abs(auc - exp["auc"]) < TOL,
              f"{auc:.4f} vs {exp['auc']}")

    # --- Gini 自洽 ---
    gini_ok = all(abs(float(mm.loc[f, "gini"]) - (2 * float(mm.loc[f, "auc"]) - 1)) < 2e-3
                  for f in GOLDEN if f in mm.index)
    check("Gini = 2×AUC-1 自洽", gini_ok)

    # --- 过拟合告警 ---
    overfit = str(mm.loc["train", "过拟合告警"]) if "过拟合告警" in mm.columns else ""
    check("过拟合告警生效（KS衰减>0.05）", overfit == "是", f"标记={overfit}")

    # --- Lift 方向正确性（防止头尾对调复发）---
    lift_ok = all(float(mm.loc[f, "lift-1%"]) > float(mm.loc[f, "lift-尾部10%"])
                  for f in GOLDEN if f in mm.index)
    check("头部Lift > 尾部Lift（方向正确）", lift_ok)

    # --- 美化生效 ---
    cf = count_conditional_formats(report_path)
    check("数据条(dataBar)已写入", cf["dataBar"] > 0, f"{cf['dataBar']} 个")
    check("色阶(colorScale)已写入", cf["colorScale"] > 0, f"{cf['colorScale']} 个")

    # --- 定版版式：分组空行 + 条件格式不跨空行 ---
    grouped = ["单变量分析", "单变量分箱按月分析", "校准度"]
    blank_detail, blank_ok = [], True
    for sn in grouped:
        empties, _, _ = sheet_layout(report_path, sn)
        if not empties:
            blank_ok = False
        blank_detail.append(f"{sn}={len(empties)}处")
    check("分组空行已插入（变量/数据集间 1 行）", blank_ok, " ".join(blank_detail))

    cross = []
    for sn in grouped:
        empties, rngs, _ = sheet_layout(report_path, sn)
        cross += [f"{sn}:{t}{r1}-{r2}" for (t, r1, r2, _) in rngs
                  if any(r1 < e < r2 for e in empties)]
    check("条件格式不跨分组空行", not cross, str(cross[:3]) if cross else "全部按组分段")

    _, rngs_m, ws_m = sheet_layout(report_path, "单变量分箱按月分析")
    month_top = [c.value for c in ws_m[2]]
    lift_bars = [1 for (t, _, _, c) in rngs_m if t == "dataBar" and month_top[c - 1] == "LIFT"]
    _, rngs_c, ws_c = sheet_layout(report_path, "校准度")
    gap_scale = [1 for (t, _, _, c) in rngs_c
                 if t == "colorScale" and ws_c.cell(row=2, column=c).value == "gap"]
    check("LIFT 数据条 + 校准度 gap 色阶", len(lift_bars) > 0 and len(gap_scale) > 0,
          f"LIFT数据条={len(lift_bars)} gap色阶={len(gap_scale)}")

    # --- 特征全景 train/test/oot 三套 ---
    fs = pd.read_excel(report_path, "特征全景", header=1)
    need = ["IV_train", "IV_test", "IV_oot", "KS_train", "KS_test", "KS_oot", "PSI_oot"]
    check("特征全景含 train/test/oot 三套", all(c in fs.columns for c in need))

    # --- 单变量分析无错位 ---
    uv = pd.read_excel(report_path, "单变量分析", header=1)
    check("单变量分析输出非空", not uv.empty, f"{len(uv)} 行")


# ===========================================================================
# B. Mock 异常分支验证
# ===========================================================================
def validate_mock() -> None:
    print("\n" + "=" * 78)
    print("B. Mock 异常分支验证（覆盖真实数据触达不到的边界）")
    print("=" * 78)

    from generate_mock_data import generate_mock_data

    paths = generate_mock_data(n=50000, seed=42,
                               out_dir=tempfile.mkdtemp(prefix="riskmockdata_"))
    df = pd.read_csv(paths["model_result"])
    wide = pd.read_csv(paths["wide"])
    score_col = "D05_mock_score"

    # 1. 单类别月（无坏样本）不崩溃
    single = df.groupby("month")["y"].agg(["count", "sum"])
    zero_bad = single[single["sum"] == 0]
    check("存在单类别月（无坏样本）场景", len(zero_bad) > 0,
          f"{list(zero_bad.index)}")
    from riskpkg.evaluation import ModelEvaluator
    ev = ModelEvaluator(y_col="y", channel_col="channle", month_col="month")
    try:
        m = ev.evaluate_by_month(df, score_col)
        ok = not m.empty
        detail = f"{len(m)} 个月，单类别月 KS={m.loc[zero_bad.index[0], 'ks']}"
    except Exception as exc:
        ok, detail = False, repr(exc)
    check("单类别月不崩溃", ok, detail)

    # 2. 类别型特征 PSI（旧版遇字符串必崩）
    try:
        psi_cat = RiskMetrics.calc_psi(wide["cat_feat"], wide["cat_feat"].sample(frac=0.5, random_state=1))
        ok = (psi_cat is not None) and not np.isnan(psi_cat)
    except Exception as exc:
        ok, psi_cat = False, repr(exc)
    check("类别型特征 PSI 可计算", ok, f"PSI={psi_cat}")

    # 3. 高缺失特征
    miss_rate = wide["high_missing_feat"].isna().mean()
    cov = RiskMetrics.calc_coverage(wide["high_missing_feat"])
    check("高缺失特征覆盖率正确", miss_rate > 0.3 and abs(cov - (1 - miss_rate)) < 1e-3,
          f"缺失率={miss_rate:.3f} 覆盖率={cov:.3f}")

    # 4. 常量特征（退化场景）IV 不崩
    try:
        iv_const = RiskMetrics.calc_iv(wide["y"], wide["const_feat"])
        ok = (iv_const is not None) and not np.isnan(iv_const)
    except Exception as exc:
        ok, iv_const = False, repr(exc)
    check("常量特征 IV 不崩溃", ok, f"IV={iv_const}")

    # 5. 漂移特征 PSI 告警
    st = StabilityMonitor()
    tr = wide[wide["data_flag"] == "train"]["drift_feat"]
    oo = wide[wide["data_flag"] == "oot"]["drift_feat"]
    psi_drift = RiskMetrics.calc_psi(tr, oo)
    check("漂移特征 PSI > 0.25 并告警", psi_drift > 0.25,
          f"PSI={psi_drift:.4f} 等级={st.classify(psi_drift)}")

    # 6. 极小分组 PSI（样本<30 应返回 NaN 而非报错）
    tiny = wide["drift_feat"].dropna().iloc[:10]
    psi_tiny = RiskMetrics.calc_psi(wide["drift_feat"].dropna(), tiny)
    check("极小分组 PSI 返回 NaN 而非报错", psi_tiny is None or np.isnan(psi_tiny),
          f"PSI={psi_tiny}")

    # 7. Mock 全流程跑通（配置直接读 mock_main_config.json，一并验证该文件可用）
    import generate_report
    cfg = json.load(open(paths["config"]))
    tmp_out = os.path.join(tempfile.mkdtemp(prefix="riskmock_"), "mock_report.xlsx")
    try:
        generate_report.get_report(cfg, score_col, paths["wide"], paths["model_result"],
                                   tmp_out, paths["importance"], top=8)
        ok = os.path.exists(tmp_out)
        detail = f"sheets={len(pd.ExcelFile(tmp_out).sheet_names)}"
    except Exception as exc:
        ok, detail = False, repr(exc)
    check("Mock 全流程报告生成", ok, detail)

    # 8. Mock 报告美化同样生效
    if ok:
        cf = count_conditional_formats(tmp_out)
        check("Mock 报告美化生效", cf["dataBar"] > 0 and cf["colorScale"] > 0,
              f"dataBar={cf['dataBar']} colorScale={cf['colorScale']}")


# ===========================================================================
def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="风控模型报告 · 端到端验证")
    ap.add_argument("--real-data", default=None,
                    help="真实数据目录（需含 main_config.json / D05_newdf.csv / "
                         "lgbdoc.csv / D05_lgb_importance.csv）；不传则跳过 A 段")
    args = ap.parse_args(argv)

    print("=" * 78)
    print("风控模型报告 · 端到端验证")
    print("=" * 78)

    try:
        validate_real(args.real_data)
    except Exception as exc:
        check("真实数据验证整体执行", False, repr(exc))

    try:
        validate_mock()
    except Exception as exc:
        check("Mock 验证整体执行", False, repr(exc))

    print("\n" + "=" * 78)
    passed = sum(1 for s, _, _ in RESULTS if s == "PASS")
    failed = len(RESULTS) - passed
    print(f"验证结果：{passed} 通过 / {failed} 失败 （共 {len(RESULTS)} 项）")
    print("=" * 78)
    for s, name, detail in RESULTS:
        if s == "FAIL":
            print(f"  FAIL: {name} {detail}")
    return 0 if failed == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
