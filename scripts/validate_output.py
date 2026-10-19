# -*- coding: utf-8 -*-
"""
产出验收（报告生成之后跑它）
================================================================================
对**已生成的 xlsx** 做版式与内容自检，回答三个问题：
  1. 该有的 sheet 在不在？空不空？
  2. 定版版式有没有被破坏？（分组空行是否被条件格式跨过）
  3. 关键指标列在不在？（KS / AUC / PSI / 校准 gap）

与 `run_validation.py`（指标数值回归）互补：那个验"数对不对"，这个验"表对不对"。

用法
------------------------------------------------------------------
python validate_output.py --report out/report.xlsx
python validate_output.py --report out/report.xlsx --expect-sheets 总览,建模效果统计
python validate_output.py --report out/report.xlsx --json
退出码：0=通过；1=存在 ERROR

分级（为什么这样分）
------------------------------------------------------------------
ERROR（判失败）：非降级类 sheet 缺失、条件格式跨空行、**空 sheet**、**关键列缺失**
WARN （仅提醒）：4 个特征类 sheet 因输入降级而缺失（属预期行为）

内容类问题（空 sheet / 关键列缺失）必须是 ERROR：它们都意味着"写入被异常吞掉"，
若只给 WARN，报告外观完整、退出码仍为 0，就会形成假绿灯。
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import warnings
from typing import Any, Dict, List, Optional, Sequence

warnings.filterwarnings("ignore")  # openpyxl 读取 xlsxwriter 扩展属性时的噪声警告

# 默认期望的 16 个 sheet（与 generate_report.DEFAULT_SHEET_ORDER 一致）
DEFAULT_EXPECT: List[str] = [
    "总览", "样本统计", "逐月样本统计", "建模样本统计", "建模效果统计",
    "逐月模型效果", "模型分10箱", "模型分20箱", "分箱坏账趋势", "PSI热力矩阵",
    "校准度", "特征全景", "特征PSI", "单变量分析", "单变量分箱按月分析", "模型图表",
]

# 关键 sheet -> 关键列（小写子串匹配，任一命中即算通过）
# 说明：实际列名是全小写且与数据集相关的，如 ks/auc、箱内坏率/区间lift、
#       psi_test/psi_oot、total_iv_train，故用子串匹配避免大小写误报。
KEY_COLUMNS: Dict[str, Sequence[str]] = {
    "建模效果统计": ("ks", "auc"),
    "模型分10箱": ("坏率", "lift"),
    "校准度": ("gap",),
    "特征PSI": ("psi",),
    "单变量分析": ("iv",),
}


def _empty_rows(ws) -> List[int]:
    """返回完全空白的行号（1-based）。"""
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if all(v is None or (isinstance(v, str) and v.strip() == "") for v in row):
            rows.append(i)
    return rows


def _crossing_formats(ws, empties: Sequence[int]) -> List[str]:
    """条件格式区间是否跨过空行（定版版式红线：不允许）。"""
    bad: List[str] = []
    try:
        ranges = list(ws.conditional_formatting)
    except Exception:
        return bad
    for cf in ranges:
        rng = getattr(cf, "sqref", None)
        if rng is None:
            continue
        for a in rng:
            m = re.match(r"[A-Z]+(\d+):[A-Z]+(\d+)", str(a))
            if not m:
                continue
            r1, r2 = int(m.group(1)), int(m.group(2))
            if any(r1 < e < r2 for e in empties):
                bad.append(str(a))
    return bad


def validate(report_path: str,
             expect_sheets: Optional[Sequence[str]] = None,
             verbose: bool = True) -> Dict[str, Any]:
    """校验一份报告，返回结构化结果。"""
    try:
        import openpyxl
    except ImportError:
        return {"ok": False, "errors": ["未安装 openpyxl，无法验收（pip install openpyxl）"]}

    errors: List[str] = []
    warnings: List[str] = []
    info: Dict[str, Any] = {}

    if not os.path.exists(report_path):
        return {"ok": False, "errors": [f"报告不存在：{report_path}"]}

    wb = openpyxl.load_workbook(report_path)
    sheets = list(wb.sheetnames)
    expect = list(expect_sheets) if expect_sheets else DEFAULT_EXPECT

    # 1. sheet 清单
    missing = [s for s in expect if s not in sheets]
    if missing:
        # 特征类 sheet 缺页多为"输入降级"（缺宽表/缺重要性），属预期行为
        degrade4 = {"特征全景", "特征PSI", "单变量分析", "单变量分箱按月分析"}
        hard_missing = [s for s in missing if s not in degrade4]
        if hard_missing:
            errors.append(f"缺少 sheet：{hard_missing}")
        if any(s in degrade4 for s in missing):
            warnings.append(
                "缺少特征类 sheet：" + str([s for s in missing if s in degrade4])
                + "（确认是否为输入降级：缺宽表 / 缺重要性 / --no-univariate）")
    info["sheets"] = sheets
    info["sheet_count"] = len(sheets)

    # 2. 空 sheet（图表页只有图片没有单元格内容，需按图片数豁免）
    empty_sheets = []
    for s in sheets:
        ws = wb[s]
        n_img = len(getattr(ws, "_images", []) or [])
        if ws.max_row <= 1 and n_img == 0:
            empty_sheets.append(s)
    if empty_sheets:
        # 空 sheet = 写入被异常吞掉后留下的壳（generate_report 的 sheet 级 except
        # 只记 WARNING，报告照出）。这类"看着完整、其实是空的"必须判失败。
        errors.append(f"疑似空 sheet（内容缺失）：{empty_sheets}")

    # 3. 版式红线：条件格式不得跨空行
    cf_cross: Dict[str, List[str]] = {}
    for s in sheets:
        ws = wb[s]
        empties = _empty_rows(ws)
        bad = _crossing_formats(ws, empties)
        if bad:
            cf_cross[s] = bad
    if cf_cross:
        errors.append(f"条件格式跨越空行（版式红线）：{cf_cross}")
    info["gap_rows"] = {s: _empty_rows(wb[s]) for s in sheets}

    # 4. 关键列
    for sheet, cols in KEY_COLUMNS.items():
        if sheet not in sheets:
            continue
        ws = wb[sheet]
        text = " ".join(
            str(v).lower()
            for row in ws.iter_rows(values_only=True) for v in row if v is not None
        )
        if not any(str(c).lower() in text for c in cols):
            errors.append(f"sheet[{sheet}] 未找到关键列 {list(cols)}（内容缺失）")

    result = {
        "ok": len(errors) == 0,
        "report": report_path,
        "errors": errors,
        "warnings": warnings,
        "info": info,
    }

    if verbose:
        print("=" * 62)
        print(f"产出验收 | {report_path}")
        print("=" * 62)
        print(f"  sheet 数：{len(sheets)}（期望 {len(expect)}）")
        print(f"  sheet：{sheets}")
        for w in warnings:
            print(f"  [WARN] {w}")
        for e in errors:
            print(f"  [ERR ] {e}")
        print("=" * 62)
        print("结论：" + ("通过" if result["ok"] else "未通过"))
    return result


def main(argv=None) -> int:
    p = argparse.ArgumentParser(description="风控模型报告 · 产出验收")
    p.add_argument("--report", required=True, help="报告 xlsx 路径")
    p.add_argument("--expect-sheets", default=None,
                   help="期望 sheet 清单（逗号分隔），默认 16 个标准 sheet")
    p.add_argument("--json", action="store_true", help="输出 JSON")
    args = p.parse_args(argv)
    expect = [s.strip() for s in args.expect_sheets.split(",")] if args.expect_sheets else None
    res = validate(args.report, expect, verbose=not args.json)
    if args.json:
        print(json.dumps(res, ensure_ascii=False, indent=2))
    return 0 if res["ok"] else 1


if __name__ == "__main__":
    sys.exit(main())
